#!/usr/bin/env python3
"""Legge il foglio "Cespiti ante 2018" dell'export xlsx e genera import_cespiti.sql (+ controllo dei totali).
Uso: python3 -I importa_cespiti.py inventario.xlsx cartella_output
Rieseguibile: svuota e ricarica cespite_storico."""
import sys, os, datetime as dt
import openpyxl

SRC, OUT = sys.argv[1], sys.argv[2]
os.makedirs(OUT, exist_ok=True)
ws = openpyxl.load_workbook(SRC, data_only=True)["Cespiti ante 2018"]

def q(v):
    if v is None or v == '': return 'NULL'
    if isinstance(v, (int, float)): return repr(round(v, 2)) if isinstance(v, float) else str(v)
    if isinstance(v, (dt.datetime, dt.date)): return "'%s'" % v.date().isoformat() if isinstance(v, dt.datetime) else "'%s'" % v.isoformat()
    return "'" + str(v).strip().replace("'", "''") + "'"

hdr_row = next(i for i, r in enumerate(ws.iter_rows(values_only=True), 1) if r and r[0] == 'Categoria')
sql = ['BEGIN;', 'DELETE FROM cespite_storico;']
tot = {'te': 0.0, 'ab': 0.0}; n = 0; dichiarati = None
for r in ws.iter_rows(min_row=hdr_row + 1, values_only=True):
    cat = r[0]
    if cat is None: continue
    if str(cat).strip().upper() == 'TOTALE':
        dichiarati = (r[8], r[9]); continue
    n += 1
    num = int(r[1])
    colh, te, ab = (None if v in (None, '') else float(v) for v in (r[7], r[8], r[9]))
    tot['te'] += te or 0; tot['ab'] += ab or 0
    sql.append("INSERT INTO cespite_storico (categoria,numero,data,fornitore,descrizione,col_h,da_trovare_eliminare,da_abbinare,note) VALUES (%s);" % ','.join(
        q(x) for x in (str(cat).strip(), num, r[2], r[3], r[4], colh, te, ab, r[10])))
sql.append('COMMIT;')
open(f'{OUT}/import_cespiti.sql', 'w').write('\n'.join(sql) + '\n')
print(f'{n} cespiti · da trovare/eliminare {tot["te"]:.2f} · da abbinare {tot["ab"]:.2f}')
if dichiarati:
    ok = abs(tot['te'] - dichiarati[0]) < 0.005 and abs(tot['ab'] - dichiarati[1]) < 0.005
    print('Totali del foglio:', dichiarati, '-> QUADRANO' if ok else '-> NON QUADRANO')
    sys.exit(0 if ok else 1)
