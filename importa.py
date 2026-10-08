#!/usr/bin/env python3
"""Legge l'export xlsx di Inventario_IT_2026 e genera import.sql + report.txt.
Uso: python3 -I importa.py inventario.xlsx cartella_output
Non modifica il foglio: tutte le correzioni stanno qui, dichiarate e ripetibili."""
import sys, re, datetime as dt, os
import openpyxl

SRC, OUT = sys.argv[1], sys.argv[2]
os.makedirs(OUT, exist_ok=True)
wb = openpyxl.load_workbook(SRC, data_only=True)
report = []
def warn(m): report.append(m)

def q(v):                       # literal SQL
    if v is None or v == '': return 'NULL'
    if isinstance(v, bool): return 'true' if v else 'false'
    if isinstance(v, (int, float)): return repr(v)
    return "'" + str(v).replace("'", "''") + "'"

def s(v):
    if v is None: return None
    v = str(v).strip()
    return v or None

def num_testo(v):
    """Numero di documento come testo: una cella numerica (4198) arriva da openpyxl come 4198.0; qui torna 4198."""
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    if isinstance(v, int) and not isinstance(v, bool): return str(v)
    return s(v)

def parse_date(v, ctx=''):
    if v is None or v == '': return None
    if isinstance(v, dt.datetime): return v.date()
    if isinstance(v, dt.date): return v
    t = str(v).strip()
    m = re.match(r'^(\d{1,2})/(\d{1,2})/(\d{2,4})$', t)
    if m:
        d, mo, y = map(int, m.groups())
        if y < 100: y += 2000
        return dt.date(y, mo, d)
    warn(f'Data non riconosciuta {ctx}: {v!r}'); return None

def parse_money(v, ctx=''):
    if v is None or v == '': return None
    if isinstance(v, (int, float)): return round(float(v), 2)
    t = re.sub(r'[€\s]', '', str(v))
    t = t.replace('.', '').replace(',', '.')
    try: return round(float(t), 2)
    except ValueError:
        warn(f'Importo non riconosciuto {ctx}: {v!r}'); return None

def rows(name, header_row=1):
    ws = wb[name]
    hdr = [s(c.value) for c in ws[header_row]]
    for r in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if not any(x not in (None, '') for x in r): continue
        yield {h: v for h, v in zip(hdr, r) if h}

# ---------- fornitori: una sola grafia ----------
ALIAS = [
    (r'effesistemi', 'Effesistemi'), (r'amazon', 'Amazon'), (r'wind', 'Wind Tre'),
    (r'media(world|market)', 'MediaWorld'), (r'groove', 'Groove to'),
    (r'realtechnology', 'Realtechnology'), (r"cart'?armata", "Cart'armata edizioni"),
    (r'ceo informatica', 'Ceo informatica'), (r'arce byte', 'Arce byte'),
    (r'multimedia', 'Multimedia'), (r'infodigita', 'Infodigita'),
    (r'tre esse', 'Tre esse'), (r'otomac', 'Otomac'), (r'juice', 'Juice'),
    (r'unieuro', 'Unieuro'), (r'setis', 'Setis'),
]
def fornitore(v):
    t = s(v)
    if not t: return None
    for pat, nome in ALIAS:
        if re.search(pat, t, re.I): return nome
    warn(f'Fornitore senza regola, tenuto com\'è: {t!r}')
    return t

# ---------- persone ----------
persone = {}      # nome completo -> id
sql = ['BEGIN;', 'SET CONSTRAINTS ALL DEFERRED;']
AZ = {"Cart'armata": 1, 'Le parole': 2}
sql.append('-- persone')
pid = 0
for r in rows('Dipendenti'):
    nome = s(r.get('Nome completo'))
    if not nome: continue
    pid += 1; persone[nome] = pid
    mail = s(r.get('Email'))
    sql.append("INSERT INTO persona (id,nome,cognome,email,azienda_id,reparto,stato,data_ingresso,data_uscita,note) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s);" % (
        pid, q(s(r.get('Nome'))), q(s(r.get('Cognome'))), q(mail.lower() if mail else None),
        q(AZ.get(s(r.get('Azienda')))), q(s(r.get('Reparto'))), q(s(r.get('Stato')) or 'Attivo'),
        q(parse_date(r.get('Data ingresso'))), q(parse_date(r.get('Data uscita'))), q(s(r.get('Note')))))
    if not mail: warn(f'Persona senza email: {nome}')
    if not s(r.get('Azienda')): warn(f'Persona senza azienda: {nome}')

def persona_id(nome, ctx):
    nome = s(nome)
    if not nome: return None
    if nome not in persone: warn(f'Persona non in Dipendenti ({ctx}): {nome}'); return None
    return persone[nome]

# ---------- asset ----------
forn, fatt, cesp = {}, {}, {}
asset_rows, asset_ids = [], {}
for r in rows('Asset'):
    cod = s(r.get('ID'))
    if not cod: continue
    aid = int(cod.split('-')[1]); asset_ids[cod] = aid
    az = AZ.get(s(r.get('Azienda')))
    if az is None: warn(f'{cod}: azienda non valida {r.get("Azienda")!r}')
    f = fornitore(r.get('Fornitore'))
    if f and f not in forn: forn[f] = len(forn) + 1
    d_acq = parse_date(r.get('Data acquisto'), cod)
    num = num_testo(r.get('N. fattura'))
    fid = None
    if num and f:
        key = (forn[f], num, d_acq)
        if key not in fatt: fatt[key] = len(fatt) + 1
        fid = fatt[key]
    elif num and not f: warn(f'{cod}: fattura {num} senza fornitore')
    cnum = r.get('N. cespite'); cid = None
    if cnum not in (None, ''):
        cnum = str(cnum).strip().upper()
        if cnum.endswith('.0'): cnum = cnum[:-2]          # numeri letti come decimali dal foglio
        if not re.fullmatch(r'[A-Z0-9][A-Z0-9._/-]{0,29}', cnum): warn(f'{cod}: cespite non valido {cnum!r}'); cnum = None
        if cnum is not None and az:
            k = (az, cnum)
            if k not in cesp: cesp[k] = len(cesp) + 1
            cid = cesp[k]
    ram = r.get('RAM (GB)'); sto = r.get('Storage (GB)')
    tip = s(r.get('Tipologia'))
    if ram not in (None, '') and tip not in ('Server',) and float(ram) > 128:
        warn(f'{cod}: RAM sospetta ({ram} GB)')
    marca = s(r.get('Marca'))
    if marca and marca.islower(): marca = marca.upper() if len(marca) <= 3 else marca.title()
    asset_rows.append((aid, cod, tip, s(r.get('Stato')), persona_id(r.get('Assegnato a'), cod), marca,
        s(r.get('Modello')), int(ram) if ram not in (None, '') else None, int(sto) if sto not in (None, '') else None,
        s(r.get('Sistema operativo')), s(r.get('Serial Number')), s(r.get('Nome dispositivo')), az, cid,
        forn.get(f) if f else None, fid, d_acq, parse_money(r.get('Importo €'), cod),
        parse_date(r.get('Data dismissione'), cod), s(r.get('Note'))))

sql.append('-- fornitori, cespiti, fatture')
for n, i in forn.items(): sql.append(f"INSERT INTO fornitore (id,nome) VALUES ({i},{q(n)});")
for (a, n), i in cesp.items(): sql.append(f"INSERT INTO cespite (id,azienda_id,numero) VALUES ({i},{a},'{n}');")
fn = {v: k for k, v in forn.items()}
for (f_id, num, d), i in fatt.items(): sql.append(f"INSERT INTO fattura (id,fornitore_id,numero,data) VALUES ({i},{f_id},{q(num)},{q(d)});")

# trigger spento durante il caricamento: i movimenti storici vengono dal foglio, non rigenerati
sql.append('ALTER TABLE asset DISABLE TRIGGER asset_movimenti;')
sql.append('-- asset')
for a in asset_rows:
    sql.append("INSERT INTO asset (id,tipologia,stato,persona_id,marca,modello,ram_gb,storage_gb,sistema_operativo,serial,hostname,azienda_id,cespite_id,fornitore_id,fattura_id,data_acquisto,importo,data_dismissione,note) VALUES (%s);" % ','.join(q(x) for x in (a[0],)+a[2:]))
sql.append('ALTER TABLE asset ENABLE TRIGGER asset_movimenti;')

# ---------- SIM ----------
by_host = {}
by_pers = {}
for a in asset_rows:
    if a[2] == 'Telefono':
        by_pers.setdefault(a[4], []).append(a[0])
        if a[11]: by_host.setdefault(a[11].upper(), []).append(a[0])
sql.append('-- sim')
for r in rows('SIM'):
    cod = s(r.get('ID'))
    if not cod: continue
    sid = int(cod.split('-')[1])
    p = persona_id(r.get('Assegnata a'), cod)
    link = None; tel = s(r.get('Telefono collegato'))
    if tel:
        c = [x for x in by_host.get(tel.upper(), []) if next(a for a in asset_rows if a[0] == x)[4] == p]
        if len(c) == 1: link = c[0]
        elif p and len(by_pers.get(p, [])) == 1: link = by_pers[p][0]
        if link is None: warn(f'{cod}: telefono "{tel}" non collegato ({r.get("Assegnata a")}) — da indicare a mano')
    sap = s(r.get('Saponetta')); note = s(r.get('Note'))
    saponetta = None
    if sap:
        if sap.lower().startswith('si'): saponetta = True
        elif sap.lower() == 'no': saponetta = False
        if sap.lower() not in ('si', 'no'): note = (note + ' · ' if note else '') + f'Saponetta: {sap}'
    num = s(r.get('Numero'))
    sql.append("INSERT INTO sim (id,numero,stato,persona_id,operatore,piano,costo_mensile,azienda_id,asset_id,giga_telefono,saponetta,giga_saponetta,note) VALUES (%s);" % ','.join(q(x) for x in (
        sid, num, s(r.get('Stato')), p, s(r.get('Operatore')), s(r.get('Piano tariffario')),
        parse_money(r.get('Costo mensile €'), cod), AZ.get(s(r.get('Azienda'))), link,
        s(r.get('Giga telefono')), saponetta, s(r.get('Giga saponetta')), note)))

# ---------- movimenti ----------
sql.append('-- movimenti')
seen = set(); dup = 0; n_mov = 0
for r in rows('Movimenti'):
    cod = s(r.get('ID asset'))
    if not cod: continue
    if cod not in asset_ids: warn(f'Movimento su asset inesistente {cod}'); continue
    d = parse_date(r.get('Data'), 'Movimenti')
    da = persona_id(r.get('Da (nome)'), 'Movimenti'); a = persona_id(r.get('A (nome)'), 'Movimenti')
    nota = s(r.get('Note')) or ''
    m = re.search(r'stato: (.+?) -> (.+)$', nota)
    prima = dopo = None
    if m:
        prima, dopo = m.group(1).strip(), m.group(2).strip()
        if prima.startswith('In magazzino'): prima = None
        if dopo.startswith('In magazzino'): dopo = None
    key = (d, s(r.get('Tipo movimento')), cod, da, a, nota)
    if key in seen: dup += 1; continue
    seen.add(key); n_mov += 1
    sql.append("INSERT INTO movimento (data,tipo,asset_id,da_persona_id,a_persona_id,stato_prima,stato_dopo,note,automatico) VALUES (%s);" % ','.join(q(x) for x in (
        d, s(r.get('Tipo movimento')), asset_ids[cod], da, a, prima, dopo, nota or None, nota.startswith('(automatico)'))))
if dup: warn(f'Movimenti doppi scartati: {dup}')

# ---------- antivirus ----------
sql.append('-- antivirus')
AV_IMPORT_IDX = len(sql)
sql.append('')   # riga riempita dopo, quando si conosce la data piu' recente del report
def av_date(v):
    if isinstance(v, dt.datetime): return v
    t = s(v)
    if not t: return None
    for f in ('%m/%d/%Y %H:%M', '%d/%m/%Y %H:%M'):
        try: return dt.datetime.strptime(t, f)           # il report esporta mm/gg/aaaa
        except ValueError: pass
    warn(f'Antivirus: data non riconosciuta {t!r}'); return None
seen_av = set(); av_max = None
for r in rows('Antivirus'):
    d = s(r.get('Dispositivo'))
    if not d or d in seen_av: continue
    seen_av.add(d)
    dd = av_date(r.get('Ultimo rilevato'))
    if dd and (av_max is None or dd > av_max): av_max = dd
    sql.append("INSERT INTO antivirus_dispositivo (import_id,dispositivo,stato,utente,sistema_operativo,ultimo_rilevato,ip_locale,mac) VALUES (1,%s,%s,%s,%s,%s,%s,%s);" % (
        q(d), q(s(r.get('Stato'))), q(s(r.get('Utente in uso'))), q(s(r.get('Sistema operativo'))),
        q(av_date(r.get('Ultimo rilevato'))), q(s(r.get('IP locale'))), q(s(r.get('MAC')))))

# l'import e' datato come il report piu' recente (non "adesso"), cosi' la soglia dei 20 giorni resta quella dell'epoca
sql[AV_IMPORT_IDX] = "INSERT INTO antivirus_import (id,importato,file_nome) VALUES (1,%s,'importazione da foglio');" % q(av_max)

# ---------- sequenze ----------
for t in ('persona', 'fornitore', 'cespite', 'fattura', 'asset', 'sim', 'movimento', 'antivirus_import'):
    sql.append(f"SELECT setval(pg_get_serial_sequence('{t}','id'), (SELECT coalesce(max(id),1) FROM {t}));")
sql.append('COMMIT;')

open(f'{OUT}/import.sql', 'w').write('\n'.join(sql) + '\n')
open(f'{OUT}/fornitori.txt', 'w').write('\n'.join(sorted(forn)) + '\n')
open(f'{OUT}/report.txt', 'w').write('\n'.join(report) + '\n')
print(f'asset {len(asset_rows)} · persone {len(persone)} · fornitori {len(forn)} · fatture {len(fatt)} · cespiti {len(cesp)} · movimenti {n_mov} · avvisi {len(report)}')
